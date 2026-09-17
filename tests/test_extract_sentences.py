import tempfile
import unittest
from pathlib import Path

from docx import Document

from scripts.extract_sentences import extract_file, extract_sentences


class ExtractSentencesTests(unittest.TestCase):
    def test_extracts_authoritative_first_source_table_in_order(self) -> None:
        document = Document()
        first = document.add_table(rows=1, cols=2)
        first.rows[0].cells[1].text = "EN"
        row = first.add_row()
        row.cells[0].text = "Вася точно дурак. \nВася точно не придет. "

        repeated = document.add_table(rows=1, cols=2)
        repeated.rows[0].cells[0].text = "model note"
        row = repeated.add_row()
        row.cells[0].text = "Вася точно дурак.\nВася точно умный."

        translations_only = document.add_table(rows=1, cols=1)
        translations_only.rows[0].cells[0].text = "EN"
        translations_only.add_row().cells[0].text = "Vasya is definitely a fool."

        self.assertEqual(
            extract_sentences(document),
            [
                "Вася точно дурак.",
                "Вася точно не придет.",
            ],
        )

    def test_applies_approved_word_order_correction(self) -> None:
        document = Document()
        table = document.add_table(rows=1, cols=1)
        table.add_row().cells[0].text = "Вася, я так и знал, ошибся."

        self.assertEqual(
            extract_sentences(document),
            ["Я так и знал, что Вася ошибся."],
        )

    def test_omits_editorial_notes_after_or_between_sentences(self) -> None:
        document = Document()
        table = document.add_table(rows=1, cols=1)
        table.add_row().cells[0].text = (
            "Вася таки не пришел. (заметка о переводе)\n"
            "Прим: это не исходное предложение\n"
            "Вася таки ошибся."
        )

        self.assertEqual(
            extract_sentences(document),
            ["Вася таки не пришел.", "Вася таки ошибся."],
        )

    def test_writes_utf8_with_one_sentence_per_line(self) -> None:
        document = Document()
        table = document.add_table(rows=1, cols=1)
        table.add_row().cells[0].text = "Вася, конечно, ошибся."

        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            source = directory / "source.docx"
            destination = directory / "sentences.txt"
            document.save(source)

            count = extract_file(source, destination)

            self.assertEqual(count, 1)
            self.assertEqual(
                destination.read_bytes(),
                "Вася, конечно, ошибся.\n".encode("utf-8"),
            )


if __name__ == "__main__":
    unittest.main()
